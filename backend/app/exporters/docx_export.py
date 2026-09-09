"""Экспорт в Word."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from ..schemas import SessionDetail, SessionStats
from .common import clock, speaker_names


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def to_docx(detail: SessionDetail, stats: SessionStats) -> bytes:
    doc = Document()

    doc.add_heading(detail.title or "Транскрипт", level=0)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = meta.add_run(
        f"{detail.created_at.strftime('%d.%m.%Y %H:%M')}  ·  "
        f"{clock(detail.duration_sec)}  ·  "
        f"говорящих: {len(detail.speakers)}  ·  "
        f"{'черновик' if detail.quality == 'draft' else 'чистовик'}"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    if stats.speakers:
        doc.add_heading("Говорящие", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for i, head in enumerate(("Имя", "Время речи", "Доля", "Слов")):
            table.rows[0].cells[i].text = head
        for row in stats.speakers:
            cells = table.add_row().cells
            cells[0].text = row.display_name
            cells[1].text = clock(row.talk_time_sec)
            cells[2].text = f"{row.share * 100:.0f}%"
            cells[3].text = str(row.word_count)

    doc.add_heading("Транскрипт", level=1)

    names = speaker_names(detail)
    colors = {sp.id: sp.color for sp in detail.speakers}
    last_speaker: int | None = None

    for seg in detail.segments:
        text = seg.text.strip()
        if not text:
            continue
        if seg.speaker_id != last_speaker:
            head = doc.add_paragraph()
            head.paragraph_format.space_before = Pt(10)
            name_run = head.add_run(f"{names.get(seg.speaker_id, 'Спикер')}  ")
            name_run.bold = True
            if seg.speaker_id in colors:
                name_run.font.color.rgb = _rgb(colors[seg.speaker_id])
            ts_run = head.add_run(clock(seg.start))
            ts_run.font.size = Pt(8)
            ts_run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
            last_speaker = seg.speaker_id
        doc.add_paragraph(text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
